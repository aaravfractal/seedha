"""Statement parsing: deterministic layouts first, Claude only as a fallback.

Layouts are recognised by their content (headers, title text), never by the
file name, because owners rename files. Known layouts are read by plain code.
Anything else goes to the StatementReader (Claude) with a strict schema; its
rows are then grounded: every amount it returns must appear in the source
text, or the row is flagged and kept out of the flip (rules 1 and 3).
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Protocol

import pdfplumber
from openpyxl import load_workbook

from app.services.checks import RawRow
from app.services.rows import ParsedRow

MAX_TEXT_FOR_LLM = 400_000
MAX_XLSX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 300
CANCEL_WORDS = frozenset({"cancelled", "canceled", "cancelled_by_guest", "cancellation", "cancel"})
DATE_FORMATS = ("%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y")
REASON_NOT_IN_SOURCE = "not_in_source"

Cell = object
Record = Mapping[str, Cell]


class ParseError(ValueError):
    """The file cannot be read as a statement; the message is owner-facing."""


class ReaderUnavailableError(RuntimeError):
    """Claude is not configured or could not read the file."""


class StatementReader(Protocol):
    """Reads an unknown statement layout into the canonical JSON shape."""

    def extract(self, text: str, filename: str) -> dict[str, object]:
        """Return {"channel": str, "rows": [...]} with amounts as printed."""


@dataclass(frozen=True, slots=True)
class ParseResult:
    """Rows from one file and how they were read."""

    layout: str
    method: str
    rows: list[ParsedRow]


# --- small converters ---------------------------------------------------------


def to_paise(value: Cell) -> int:
    """A printed or numeric rupee amount as paise: '4,200.00' -> 420000.

    Accepts a leading minus or (brackets) for negatives, and Rs/INR/₹ marks.
    """
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        raise ParseError(f"Not an amount: {value!r}")
    text = str(value).strip()
    negative = text.startswith("-") or (text.startswith("(") and text.endswith(")"))
    cleaned = re.sub(r"[\s,₹()]|INR|Rs\.?", "", text).lstrip("-")
    try:
        amount = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ParseError(f"Not an amount: {text!r}") from exc
    paise = amount * 100
    if paise != paise.to_integral_value():
        raise ParseError(f"More than two decimals: {text!r}")
    return -int(paise) if negative else int(paise)


def to_date(value: Cell) -> date:
    """A statement date in any of the formats Indian exports use."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()  # noqa: DTZ007  (calendar date only)
        except ValueError:
            continue
    raise ParseError(f"Not a date: {text!r}")


def to_int(value: Cell) -> int:
    """A whole number such as nights or a room number."""
    try:
        return int(str(value).strip())
    except ValueError as exc:
        raise ParseError(f"Not a whole number: {value!r}") from exc


def to_status(value: Cell) -> str:
    """Normalise every channel's wording to 'stay' or 'cancelled'."""
    return "cancelled" if str(value).strip().lower() in CANCEL_WORDS else "stay"


def text_or_none(value: Cell) -> str | None:
    """A trimmed string, or None for blank cells."""
    text = str(value).strip() if value is not None else ""
    return text or None


def money_row(status: str, list_amount: Cell, payout: Cell, **deductions: Cell) -> RawRow:
    """Money lines in paise. Deductions are keyed discount, commission, gst,
    fees, tds, tcs (missing means zero). Every one must be non-negative."""
    unknown = set(deductions) - {"discount", "commission", "gst", "fees", "tds", "tcs"}
    if unknown:
        raise ParseError(f"Unknown money columns {sorted(unknown)}")
    lines = {
        "discount": to_paise(deductions.get("discount", 0)),
        "commission": to_paise(deductions.get("commission", 0)),
        "gst_on_commission": to_paise(deductions.get("gst", 0)),
        "fees": to_paise(deductions.get("fees", 0)),
        "tds": to_paise(deductions.get("tds", 0)),
        "gst_tcs": to_paise(deductions.get("tcs", 0)),
    }
    gross = to_paise(list_amount)
    if gross < 0 or any(v < 0 for v in lines.values()):
        raise ParseError("A list amount or deduction is negative")
    return RawRow(status=status, list_amount=gross, payout_reported=to_paise(payout), **lines)


# --- file type ----------------------------------------------------------------


def sniff_type(filename: str, data: bytes) -> str:
    """'pdf', 'xlsx' or 'csv', checked against the bytes, not just the name."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if data.startswith(b"%PDF"):
        if suffix != "pdf":
            raise ParseError("This file is a PDF but is not named .pdf")
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        if suffix not in {"xlsx", "xlsm"}:
            raise ParseError("This file is a zip/Excel file but is not named .xlsx")
        return "xlsx"
    if suffix == "csv":
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ParseError("This CSV is not UTF-8 text") from exc
        return "csv"
    raise ParseError("Upload a PDF, Excel (.xlsx) or CSV statement")


# --- tabular readers ----------------------------------------------------------


def records(header: Sequence[Cell], rows: Iterable[Sequence[Cell]]) -> Iterator[dict[str, Cell]]:
    """Header-keyed records, skipping blank lines and repeated header rows."""
    names = [str(h).strip() if h is not None else "" for h in header]
    for row in rows:
        values = list(row)
        if not any(v not in (None, "") for v in values):
            continue
        if [str(v).strip() if v is not None else "" for v in values] == names:
            continue
        yield dict(zip(names, values, strict=False))


def xlsx_table(data: bytes) -> tuple[list[Cell], list[list[Cell]]]:
    """Header and rows of the first sheet. Refuses zip bombs before unpacking."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ParseError("This Excel file could not be opened") from exc
    if unpacked > MAX_XLSX_UNPACKED_BYTES:
        raise ParseError("This Excel file is too large once unpacked")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ParseError("This Excel file could not be opened") from exc
    sheet = workbook.worksheets[0]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        raise ParseError("The first sheet is empty")
    return rows[0], rows[1:]


def csv_table(data: bytes) -> tuple[list[Cell], list[list[Cell]]]:
    """Header and rows of a UTF-8 CSV."""
    reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
    rows: list[list[Cell]] = [list(r) for r in reader]
    if not rows:
        raise ParseError("The CSV is empty")
    return rows[0], rows[1:]


def pdf_tables(data: bytes) -> tuple[str, list[list[Cell]]]:
    """All page text and all table rows of a PDF."""
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise ParseError(f"PDFs can have at most {MAX_PDF_PAGES} pages. Split it by month.")
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            table_rows: list[list[Cell]] = []
            for page in pdf.pages:
                for table in page.extract_tables():
                    table_rows.extend(list(r) for r in table)
    except ParseError:
        raise
    except Exception as exc:  # pdfminer raises many unrelated types on bad input
        raise ParseError("This PDF could not be opened") from exc
    return text, table_rows


# --- known layouts ------------------------------------------------------------


def mmt_row(r: Record) -> ParsedRow:
    """MakeMyTrip payout PDF line."""
    status = to_status(r["Status"])
    return ParsedRow(
        channel="mmt",
        external_ref=str(r["Booking ID"]).strip(),
        status=status,
        check_in=to_date(r["Check-in"]),
        nights=to_int(r["Nights"]),
        guest_name=str(r["Guest"]).strip(),
        guest_city=text_or_none(r["City"]),
        money=money_row(
            list_amount=r["Room charges"],
            discount=r["Hotel promo"],
            commission=r["Commission"],
            gst=r["GST on commission"],
            tds=r["TDS"],
            tcs=r["TCS"],
            payout=r["Net payable"],
            status=status,
        ),  # fmt: skip
    )


def bcom_row(r: Record) -> ParsedRow:
    """Booking.com payouts sheet line; nights from arrival and departure."""
    status = to_status(r["Status"])
    arrival = to_date(r["Arrival"])
    return ParsedRow(
        channel="bcom",
        external_ref=f"BK{to_int(r['Reservation number'])}",
        status=status,
        check_in=arrival,
        nights=(to_date(r["Departure"]) - arrival).days,
        guest_name=str(r["Guest name"]).strip(),
        guest_city=text_or_none(r["Guest city"]),
        money=money_row(
            list_amount=r["Room price"],
            discount=r["Promotion discount"],
            commission=r["Commission"],
            gst=r["GST on commission"],
            tds=r["TDS"],
            payout=r["Payout amount"],
            status=status,
        ),  # fmt: skip
    )


def agoda_row(r: Record) -> ParsedRow:
    """Agoda statement CSV line; nights from check-in and check-out."""
    status = to_status(r["Booking Status"])
    arrival = to_date(r["Check-in Date"])
    return ParsedRow(
        channel="agoda",
        external_ref=f"AG{to_int(r['Booking ID'])}",
        status=status,
        check_in=arrival,
        nights=(to_date(r["Check-out Date"]) - arrival).days,
        guest_name=str(r["Guest Name"]).strip(),
        guest_city=text_or_none(r["Guest City"]),
        money=money_row(
            list_amount=r["Selling Rate (INR)"],
            discount=r["Hotel Promotion (INR)"],
            commission=r["Commission (INR)"],
            gst=r["GST on Commission (INR)"],
            tds=r["TDS (INR)"],
            payout=r["Net Amount (INR)"],
            status=status,
        ),  # fmt: skip
    )


def ta_row(r: Record) -> ParsedRow:
    """Travel agent register line (commission without GST, D-001)."""
    status = to_status(r["Status"])
    return ParsedRow(
        channel="ta",
        external_ref=str(r["Voucher no"]).strip(),
        status=status,
        check_in=to_date(r["Date"]),
        nights=to_int(r["Nights"]),
        guest_name=str(r["Guest"]).strip(),
        guest_city=text_or_none(r["City"]),
        room=to_int(r["Room"]) if r.get("Room") not in (None, "") else None,
        money=money_row(
            list_amount=r["Tariff"],
            commission=r["Commission"],
            payout=r["Amount received"],
            status=status,
        ),  # fmt: skip
    )


def register_row(r: Record) -> ParsedRow:
    """Hotel register line: walk-in (cash) or direct (UPI fee)."""
    source = str(r["Source"]).strip().lower()
    if source not in {"walk-in", "walkin", "direct"}:
        raise ParseError(f"Unknown register source {r['Source']!r}")
    return ParsedRow(
        channel="direct" if source == "direct" else "walkin",
        external_ref=str(r["Bill no"]).strip(),
        status="stay",
        check_in=to_date(r["Date"]),
        nights=to_int(r["Nights"]),
        guest_name=str(r["Guest"]).strip(),
        guest_city=text_or_none(r["City"]),
        room=to_int(r["Room"]),
        money=money_row(
            list_amount=r["Tariff"], fees=r["UPI fee"], payout=r["Amount received"], status="stay"
        ),  # fmt: skip
    )


@dataclass(frozen=True, slots=True)
class Layout:
    """A known statement layout: the headers that identify it and its row reader."""

    name: str
    kind: str
    required: frozenset[str]
    reader: Callable[[Record], ParsedRow]


LAYOUTS: tuple[Layout, ...] = (
    Layout("makemytrip-payout-pdf", "pdf",
           frozenset({"Booking ID", "Room charges", "Hotel promo", "Net payable"}), mmt_row),
    Layout("booking-com-payouts", "xlsx",
           frozenset({"Reservation number", "Room price", "Payout amount"}), bcom_row),
    Layout("agoda-statement", "csv",
           frozenset({"Booking ID", "Selling Rate (INR)", "Net Amount (INR)"}), agoda_row),
    Layout("travel-agent-register", "xlsx",
           frozenset({"Voucher no", "Agent", "Tariff", "Amount received"}), ta_row),
    Layout("hotel-register", "xlsx",
           frozenset({"Bill no", "Source", "Tariff", "UPI fee", "Amount received"}),
           register_row),
)  # fmt: skip


def match_layout(kind: str, header: Sequence[Cell]) -> Layout | None:
    """The known layout whose identifying headers are all present."""
    names = {str(h).strip() for h in header if h is not None}
    for layout in LAYOUTS:
        if layout.kind == kind and layout.required <= names:
            return layout
    return None


def read_layout(
    layout: Layout, header: Sequence[Cell], rows: Iterable[Sequence[Cell]]
) -> list[ParsedRow]:
    """Apply a layout to every data row; say which line failed on error."""
    parsed: list[ParsedRow] = []
    for line_no, record in enumerate(records(header, rows), start=2):
        try:
            parsed.append(layout.reader(record))
        except (KeyError, ParseError) as exc:
            raise ParseError(f"Line {line_no}: {exc}") from exc
    return parsed


# --- Claude fallback ----------------------------------------------------------

CHANNEL_ALIASES = {
    "makemytrip": "mmt", "mmt": "mmt", "goibibo": "mmt", "booking.com": "bcom",
    "booking": "bcom", "agoda": "agoda", "travel agent": "ta", "travel agents": "ta",
    "agent": "ta", "walk-in": "walkin", "walkin": "walkin", "direct": "direct",
}  # fmt: skip
AMOUNT_FIELDS = ("list_amount", "discount", "commission", "gst_on_commission", "fees", "tds",
                 "tcs", "payout")  # fmt: skip


def normalise_amount_text(text: str) -> str:
    """Digits and sign only, trailing decimal zeros dropped: '4,200.50' -> '4200.5'."""
    digits = re.sub(r"[^\d.\-]", "", text)
    if "." in digits:
        digits = digits.rstrip("0").rstrip(".")
    return digits


GROUPED_NUMBER = re.compile(r"-?\d{1,3}(?:,\d{2,3})+(?:\.\d+)?|-?\d+(?:\.\d+)?")


def source_amounts(text: str) -> set[str]:
    """Every number printed in free text (PDF), normalised for grounding.

    Understands Indian (1,23,456) and international (123,456) grouping.
    """
    return {normalise_amount_text(m) for m in GROUPED_NUMBER.findall(text)}


def cell_amounts(header: Sequence[Cell], rows: Iterable[Sequence[Cell]]) -> set[str]:
    """Every cell of a table, normalised, so CSV commas never merge cells."""
    printed: set[str] = set()
    for row in [header, *rows]:
        for cell in row:
            if cell is not None:
                printed |= source_amounts(str(cell))
                printed.add(normalise_amount_text(str(cell)))
    return printed


def grounded(row: Mapping[str, object], printed: set[str]) -> bool:
    """True when every non-zero amount the reader returned is in the source."""
    for field in AMOUNT_FIELDS:
        value = normalise_amount_text(str(row.get(field, "0")))
        if value not in {"", "0", "-0"} and value.lstrip("-") not in {
            p.lstrip("-") for p in printed
        }:
            return False
    return True


def llm_row(raw: Mapping[str, object], channel: str, printed: set[str]) -> ParsedRow:
    """One Claude-extracted row, checked and grounded."""
    status = "cancelled" if str(raw.get("status", "")).lower() == "cancelled" else "stay"
    row = ParsedRow(
        channel=channel,
        external_ref=str(raw["ref"]).strip(),
        status=status,
        check_in=to_date(raw["check_in"]),
        nights=to_int(raw["nights"]),
        guest_name=str(raw["guest_name"]).strip(),
        guest_city=text_or_none(raw.get("guest_city")),
        money=money_row(
            list_amount=raw["list_amount"],
            discount=raw.get("discount", 0),
            commission=raw.get("commission", 0),
            gst=raw.get("gst_on_commission", 0),
            fees=raw.get("fees", 0),
            tds=raw.get("tds", 0),
            tcs=raw.get("tcs", 0),
            payout=raw["payout"],
            status=status,
        ),  # fmt: skip
    )
    if not grounded(raw, printed):
        return row.flagged(REASON_NOT_IN_SOURCE, "An amount read by Claude is not in the file")
    return row


def read_with_llm(
    text: str, filename: str, reader: StatementReader | None, printed: set[str]
) -> ParseResult:
    """Send an unknown layout to Claude and ground what comes back."""
    if reader is None:
        raise ReaderUnavailableError(
            "This layout is new to Seedha and needs Claude to read it. "
            "Claude is not configured (simulated mode), so add ANTHROPIC_API_KEY or "
            "export the statement as CSV."
        )
    if len(text) > MAX_TEXT_FOR_LLM:
        raise ParseError("This file is too large to read in one go. Split it by month.")
    payload = reader.extract(text, filename)
    channel = CHANNEL_ALIASES.get(str(payload.get("channel", "")).strip().lower())
    if channel is None:
        raise ParseError("Could not tell which channel this statement is from")
    raw_rows = payload.get("rows")
    if not isinstance(raw_rows, list) or not raw_rows:
        raise ParseError("No booking lines were found in this file")
    rows: list[ParsedRow] = []
    for index, raw in enumerate(raw_rows, start=1):
        if not isinstance(raw, Mapping):
            raise ParseError(f"Line {index}: unreadable")
        try:
            rows.append(llm_row(raw, channel, printed))
        except (KeyError, ParseError) as exc:
            raise ParseError(f"Line {index}: {exc}") from exc
    return ParseResult(layout="unknown", method="claude", rows=rows)


def table_text(header: Sequence[Cell], rows: Iterable[Sequence[Cell]]) -> str:
    """A table as CSV text for Claude."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["" if h is None else h for h in header])
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return buffer.getvalue()


def parse_statement(filename: str, data: bytes, reader: StatementReader | None) -> ParseResult:
    """Read a statement: known layout by plain code, else Claude, else error."""
    kind = sniff_type(filename, data)
    if kind == "pdf":
        text, table = pdf_tables(data)
        header = table[0] if table else []
        layout = match_layout("pdf", header)
        if layout is not None:
            return ParseResult(layout.name, "layout", read_layout(layout, header, table[1:]))
        return read_with_llm(text, filename, reader, source_amounts(text))
    header, rows = xlsx_table(data) if kind == "xlsx" else csv_table(data)
    layout = match_layout(kind, header)
    if layout is not None:
        return ParseResult(layout.name, "layout", read_layout(layout, header, rows))
    return read_with_llm(table_text(header, rows), filename, reader, cell_amounts(header, rows))
